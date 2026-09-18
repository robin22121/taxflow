"""프로덕션 DB에 샘플 계정/거래처/급여 이력 seed.

- 2798102355 계정 (이지원천 데모사무소, 홍길동): 기존 데이터 정리 후 샘플 5개 거래처 + 1~7월 급여 이력
- 1234567890 계정 (테스트 세무사사무소, 홍길동, 비어있음)
- 다른 사무소의 데이터는 건드리지 않음

실행 (vm-node):
    sudo -u ubuntu bash -c 'cd /opt/taxflow/backend && /opt/taxflow/venv/bin/python -m app.scripts.seed_prod_samples'
"""

from __future__ import annotations

import asyncio
import logging
import re
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import delete, select

from app.core.security import hash_password
from app.db import SessionLocal
from app.models import (
    Client,
    CollectionEvent,
    CollectionSession,
    CollectionSessionStatus,
    Employee,
    ClientFilingResult,
    EmploymentStatus,
    MonthlyFiling,
    MonthlyFilingStatus,
    PayrollEntry,
    SecureToken,
    TaxOffice,
    User,
)
from app.models.employee_change import EmployeeChangeRequest
from app.models.payroll import IncomeType, MatchStatus
from app.models.tax_office import CustomerClass, OfficeApprovalStatus
from app.services.crypto import encrypt_rrn

logging.basicConfig(level=logging.INFO, format="%(message)s")
log = logging.getLogger("seed_prod")

# samples/는 /opt/taxflow/samples 로 업로드 됨
SAMPLES = Path(__file__).resolve().parents[3] / "samples"

PASSWORD = "admin1234!"
PERIODS = [f"2026-{m:02d}" for m in range(1, 8)]


def parse_hired(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        try:
            return date.fromisoformat(v)
        except ValueError:
            return None
    return None


def rrn_last4(raw: str) -> str:
    return re.sub(r"[\s-]", "", raw or "")[-4:]


def read_clients_master() -> list[dict]:
    wb = load_workbook(SAMPLES / "clients.xlsx", data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    idx = {h: i for i, h in enumerate(header)}
    return [{
        "biz_number": r[idx["사업자번호"]],
        "name": r[idx["상호"]],
        "industry": r[idx["업종"]],
        "is_corp": r[idx["법인/개인"]] == "법인",
    } for r in rows[1:]]


def read_employees(biz_number: str) -> list[dict]:
    wb = load_workbook(SAMPLES / biz_number / "employees.xlsx", data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    idx = {h: i for i, h in enumerate(rows[0])}
    return [{
        "code": r[idx["사원코드"]],
        "name": r[idx["사원명"]],
        "rrn": str(r[idx["주민번호"]] or "").strip(),
        "hired": parse_hired(r[idx["입사일"]]),
        "department": r[idx["부서"]],
        "position": r[idx["직급"]],
        "job_type": r[idx["직종"]],
        "income_type": r[idx["소득구분"]] or "WAGE",
        "base_salary": int(r[idx["기본급"]] or 0),
        "meal": int(r[idx["식대(비과세)"]] or 0),
        "car": int(r[idx["자가운전(비과세)"]] or 0),
        "childcare": int(r[idx["육아수당(비과세)"]] or 0),
    } for r in rows[1:]]


def read_payroll(biz_number: str, client_name: str, period: str) -> list[dict]:
    path = SAMPLES / biz_number / f"{client_name}-{period.replace('-', '')}.xlsx"
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    out = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 2 or row[0] == "합계" or row[0] is None:
            continue
        out.append({
            "code": str(row[0]),
            "name": (row[1] or "").strip() if isinstance(row[1], str) else (row[1] or ""),
            "department": row[2] or "",
            "position": row[3] or "",
            "job_type": row[4] or "",
            "base": int(row[5] or 0),
            "bonus": int(row[6] or 0),
            "meal": int(row[7] or 0),
            "car": int(row[8] or 0),
            "childcare": int(row[9] or 0),
            "total_pay": int(row[10] or 0),
            "np": int(row[11] or 0),
            "hi": int(row[12] or 0),
            "ei": int(row[13] or 0),
            "ltc": int(row[14] or 0),
            "income_tax": int(row[15] or 0),
            "local_tax": int(row[16] or 0),
            "student_loan": int(row[17] or 0),
            "settlement_ins": int(row[18] or 0),
            "rent_support": int(row[19] or 0),
            "total_ded": int(row[20] or 0),
            "net": int(row[21] or 0),
        })
    return out


async def ensure_office(db, biz_number: str, name: str) -> TaxOffice:
    office = (await db.execute(
        select(TaxOffice).where(TaxOffice.business_number == biz_number)
    )).scalars().first()
    if office is None:
        office = TaxOffice(
            name=name,
            short_code=f"EZOC-{biz_number[-4:]}",
            business_number=biz_number,
            representative="홍길동",
            phone="010-3640-5642",
            email="robin1q84@gmail.com",
            approval_status=OfficeApprovalStatus.APPROVED,
            approved_at=datetime.now(timezone.utc),
            customer_class=CustomerClass.REGULAR,
        )
        db.add(office)
        await db.flush()
        log.info("[office] 신규 %s (%s)", name, biz_number)
    else:
        office.name = name
        office.representative = "홍길동"
        office.phone = "010-3640-5642"
        office.email = "robin1q84@gmail.com"
        office.approval_status = OfficeApprovalStatus.APPROVED
        log.info("[office] 갱신 %s (%s)", name, biz_number)
    return office


async def ensure_user(db, login_id: str, office_id: str) -> User:
    user = (await db.execute(
        select(User).where(User.email == login_id)
    )).scalars().first()
    if user is None:
        user = User(
            tax_office_id=office_id,
            email=login_id,
            password_hash=hash_password(PASSWORD),
            name="홍길동",
            is_admin=True,
        )
        db.add(user)
        log.info("[user]   신규 %s / %s", login_id, PASSWORD)
    else:
        user.tax_office_id = office_id
        user.password_hash = hash_password(PASSWORD)
        user.name = "홍길동"
        log.info("[user]   갱신 %s / %s", login_id, PASSWORD)
    await db.flush()
    return user


async def wipe_office_data(db, office: TaxOffice) -> None:
    """해당 사무소에 속한 clients/employees/filings/payroll 등 전부 삭제."""
    clients = (await db.execute(
        select(Client).where(Client.tax_office_id == office.id)
    )).scalars().all()
    client_ids = [c.id for c in clients]
    filings = (await db.execute(
        select(MonthlyFiling).where(MonthlyFiling.tax_office_id == office.id)
    )).scalars().all()
    filing_ids = [f.id for f in filings]

    if filing_ids:
        await db.execute(delete(PayrollEntry).where(PayrollEntry.monthly_filing_id.in_(filing_ids)))
        await db.execute(delete(CollectionEvent).where(CollectionEvent.session_id.in_(
            select(CollectionSession.id).where(CollectionSession.monthly_filing_id.in_(filing_ids))
        )))
        await db.execute(delete(SecureToken).where(SecureToken.collection_session_id.in_(
            select(CollectionSession.id).where(CollectionSession.monthly_filing_id.in_(filing_ids))
        )))
        await db.execute(delete(CollectionSession).where(CollectionSession.monthly_filing_id.in_(filing_ids)))
        await db.execute(delete(MonthlyFiling).where(MonthlyFiling.id.in_(filing_ids)))
    if client_ids:
        # client·employee 참조하는 테이블들 먼저 삭제
        await db.execute(delete(SecureToken).where(SecureToken.client_id.in_(client_ids)))
        await db.execute(delete(ClientFilingResult).where(ClientFilingResult.client_id.in_(client_ids)))
        await db.execute(delete(EmployeeChangeRequest).where(EmployeeChangeRequest.client_id.in_(client_ids)))
        from app.models.client_payroll_default import ClientPayrollDefault
        await db.execute(delete(ClientPayrollDefault).where(ClientPayrollDefault.client_id.in_(client_ids)))
        await db.execute(delete(Employee).where(Employee.client_id.in_(client_ids)))
        await db.execute(delete(Client).where(Client.id.in_(client_ids)))
    await db.flush()
    log.info("[wipe]   %s: clients=%d filings=%d 삭제", office.name, len(client_ids), len(filing_ids))


async def main() -> None:
    async with SessionLocal() as db:
        # ── 1. 2798102355 사무소 준비 (기존 데이터 정리 후 샘플 이식) ──
        office_main = await ensure_office(db, "2798102355", "이지원천 데모사무소")
        await ensure_user(db, "2798102355", office_main.id)
        await wipe_office_data(db, office_main)

        # ── 2. 1234567890 빈 계정 준비 ──
        office_test = await ensure_office(db, "1234567890", "테스트 세무사사무소")
        await ensure_user(db, "1234567890", office_test.id)

        # ── 3. 5개 샘플 거래처 + 직원 + 급여 seed (office_main 밑) ──
        clients_master = read_clients_master()
        client_by_biz: dict[str, Client] = {}
        for cm in clients_master:
            c = Client(
                tax_office_id=office_main.id,
                business_name=cm["name"],
                business_number=cm["biz_number"],
                representative="홍길동",
                contact_phone="010-3640-5642",
                contact_email="robin1q84@gmail.com",
                is_corporation=cm["is_corp"],
            )
            db.add(c)
            client_by_biz[cm["biz_number"]] = c
        await db.flush()
        log.info("[seed]   거래처 %d개 생성", len(client_by_biz))

        # Employees
        emp_lookup: dict[str, dict[str, Employee]] = {}
        for cm in clients_master:
            biz = cm["biz_number"]
            client = client_by_biz[biz]
            emp_lookup[biz] = {}
            for e in read_employees(biz):
                try:
                    rrn_enc = encrypt_rrn(e["rrn"]) if e["rrn"] else None
                except Exception:
                    rrn_enc = None
                emp = Employee(
                    client_id=client.id,
                    name=e["name"],
                    rrn_encrypted=rrn_enc,
                    rrn_last4=rrn_last4(e["rrn"]) if e["rrn"] else None,
                    employee_code=e["code"],
                    department=e["department"],
                    position=e["position"],
                    job_type=e["job_type"],
                    hired_at=e["hired"],
                    status=EmploymentStatus.ACTIVE,
                )
                db.add(emp)
                emp_lookup[biz][e["code"]] = emp
        await db.flush()
        total_emp = sum(len(v) for v in emp_lookup.values())
        log.info("[seed]   직원 %d명 생성", total_emp)

        # Monthly filings + collection sessions + payroll (2026-01~07)
        total_entries = 0
        for period in PERIODS:
            filing = MonthlyFiling(
                tax_office_id=office_main.id,
                period=period,
                status=MonthlyFilingStatus.APPROVED,
            )
            db.add(filing)
            await db.flush()

            period_entries = 0
            for cm in clients_master:
                client = client_by_biz[cm["biz_number"]]
                session = CollectionSession(
                    monthly_filing_id=filing.id,
                    client_id=client.id,
                    request_token=f"prod-sample-{period}-{cm['biz_number']}",
                    status=CollectionSessionStatus.APPROVED,
                )
                db.add(session)
                await db.flush()

                rows = read_payroll(cm["biz_number"], cm["name"], period)
                pay_day = date(int(period[:4]), int(period[5:]), 25)
                for r in rows:
                    emp = emp_lookup[cm["biz_number"]].get(r["code"])
                    non_taxable = r["meal"] + r["car"] + r["childcare"]
                    taxable = r["total_pay"] - non_taxable
                    a_code = "A01" if client.is_corporation else "A02"
                    db.add(PayrollEntry(
                        monthly_filing_id=filing.id,
                        collection_session_id=session.id,
                        client_id=client.id,
                        employee_id=emp.id if emp else None,
                        raw_name=r["name"] or r["code"],
                        income_type=IncomeType.WAGE,
                        a_code=a_code,
                        total_amount=r["total_pay"],
                        salary_amount=r["base"],
                        bonus_amount=r["bonus"],
                        non_taxable=non_taxable,
                        meal_amount=r["meal"],
                        car_amount=r["car"],
                        childcare_amount=r["childcare"],
                        taxable=taxable,
                        national_pension=r["np"],
                        health_insurance=r["hi"],
                        employment_insurance=r["ei"],
                        longterm_care=r["ltc"],
                        income_tax=r["income_tax"],
                        local_tax=r["local_tax"],
                        student_loan=r["student_loan"],
                        settlement_insurance=r["settlement_ins"],
                        rent_support=r["rent_support"],
                        payment_date=pay_day,
                        match_status=MatchStatus.MATCHED if emp else MatchStatus.NEW_HIRE_SUSPECTED,
                        approved=True,
                        dependents=1,
                    ))
                    period_entries += 1

            await db.flush()
            filing.total_entries = period_entries
            filing.total_clients = len(clients_master)
            total_entries += period_entries
            log.info("[seed]   %s: 항목 %d건", period, period_entries)

        await db.commit()
        log.info("\n완료. 총 급여 %d건.", total_entries)
        log.info("로그인:")
        log.info("  2798102355 / %s  (이지원천 데모사무소, 5거래처)", PASSWORD)
        log.info("  1234567890 / %s  (테스트 세무사사무소, 0거래처)", PASSWORD)


if __name__ == "__main__":
    asyncio.run(main())
