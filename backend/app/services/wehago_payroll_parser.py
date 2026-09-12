"""위하고T 22컬럼 급여대장 결정론적 파서.

LLM에 넘기지 않고 openpyxl 로 직접 컬럼을 매핑한다.

배경: ``file_intake._excel_to_text`` 는 2행 병합 헤더(1행 수당/공제 그룹, 2행 16개 항목)를
탭 텍스트로 flatten 하여 LLM 파서에 넘긴다. LLM 이 케이스마다 다른 방식으로 헤더-값 매핑에
실패해 4대보험 값이 왜곡되는 사례가 확인됨((주)창원 2026-08 사건). 위하고T 22컬럼 스키마는
``payroll_excel.ITEM_HEADERS`` 로 완전히 고정돼 있으므로, 시그니처가 일치하면 결정론적으로
직접 읽는 것이 안정적이다.

시그니처 불일치 시 ``None`` 을 반환해 호출자가 기존 LLM 경로로 폴백하게 한다.
"""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook

logger = logging.getLogger(__name__)


# ── 헤더 시그니처 (payroll_excel.py 기준 22컬럼) ──
_ROW1_A_TO_E = ("사원코드", "사원명", "부서", "직급", "직종")
_ROW1_F = "수당"
_ROW1_L = "공제"
_ROW1_V = "차인지급액"
_ROW2_HEADERS = (
    "기본급", "상여", "식대", "자가운전", "육아", "지급액계",
    "국민연금", "건강보험", "고용보험", "장기요양보험료",
    "소득세", "지방소득세", "학자금상환액", "정산보험료", "월세지원금", "공제액계",
)


@dataclass(slots=True)
class WehagoPayrollRow:
    """위하고T 22컬럼 급여대장의 데이터 행 하나."""
    employee_code: str          # A
    name: str                    # B
    department: str              # C
    position: str                # D
    job_type: str                # E
    base_salary: int             # F 기본급
    bonus: int                   # G 상여
    meal_amount: int             # H 식대
    car_amount: int              # I 자가운전
    childcare_amount: int        # J 육아
    total_amount: int            # K 지급액계
    national_pension: int        # L 국민연금
    health_insurance: int        # M 건강보험
    employment_insurance: int    # N 고용보험
    longterm_care: int           # O 장기요양보험료
    income_tax: int              # P 소득세
    local_tax: int               # Q 지방소득세
    student_loan: int            # R 학자금상환액
    settlement_insurance: int    # S 정산보험료
    rent_support: int            # T 월세지원금
    deduction_total: int         # U 공제액계
    net_payment: int             # V 차인지급액


def _norm(v: Any) -> str:
    """헤더 셀 문자열을 비교 가능한 형태로 정규화."""
    if v is None:
        return ""
    return str(v).strip().replace(" ", "").replace("　", "")


def _to_int(v: Any) -> int:
    """숫자 셀을 int 로. 실측 원본은 쉼표 포함 문자열이 섞여 있어 정규화한다."""
    if v is None or v == "":
        return 0
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, (int, float)):
        return int(v)
    s = str(v).replace(",", "").strip()
    if not s:
        return 0
    try:
        return int(float(s))
    except (TypeError, ValueError):
        return 0


def _row1_matches(row1: list[Any]) -> bool:
    if len(row1) < 22:
        return False
    if any(_norm(row1[i]) != h for i, h in enumerate(_ROW1_A_TO_E)):
        return False
    return (
        _norm(row1[5]) == _ROW1_F
        and _norm(row1[11]) == _ROW1_L
        and _norm(row1[21]) == _ROW1_V
    )


def _row2_matches(row2: list[Any]) -> bool:
    # F(index 5)~U(index 20) 16개 항목
    if len(row2) < 21:
        return False
    for i, expected in enumerate(_ROW2_HEADERS):
        actual = _norm(row2[5 + i])
        # "자가운전" vs "자가운전보조금" 처럼 접미사 확장 허용
        if actual != expected and not actual.startswith(expected):
            return False
    return True


def parse_wehago_workbook(blob: bytes) -> list[WehagoPayrollRow] | None:
    """22컬럼 위하고T 급여대장이면 결정론적으로 파싱한다.

    스키마가 다르면 ``None`` — 호출자는 LLM 파서로 폴백해야 한다.
    합계 행(A열 = "합계" 또는 빈 사원코드)과 완전 빈 행은 스킵한다.
    """
    try:
        wb = load_workbook(io.BytesIO(blob), data_only=True, read_only=True)
    except Exception as e:  # noqa: BLE001 — openpyxl 은 다양한 예외를 던짐
        logger.debug("[wehago-parser] cannot load workbook: %s", e)
        return None

    for ws in wb.worksheets:
        rows_iter = ws.iter_rows(values_only=True)
        try:
            row1 = list(next(rows_iter))
            row2 = list(next(rows_iter))
        except StopIteration:
            continue
        if not _row1_matches(row1) or not _row2_matches(row2):
            continue

        parsed: list[WehagoPayrollRow] = []
        for raw in rows_iter:
            if raw is None:
                continue
            cells = list(raw) + [None] * (22 - len(raw))
            code = _norm(cells[0])
            # 합계 행·빈 행 스킵
            if not code or code == "합계":
                continue
            name = "" if cells[1] is None else str(cells[1]).strip()
            if not name:
                continue
            parsed.append(WehagoPayrollRow(
                employee_code=code,
                name=name,
                department=str(cells[2] or "").strip(),
                position=str(cells[3] or "").strip(),
                job_type=str(cells[4] or "").strip(),
                base_salary=_to_int(cells[5]),
                bonus=_to_int(cells[6]),
                meal_amount=_to_int(cells[7]),
                car_amount=_to_int(cells[8]),
                childcare_amount=_to_int(cells[9]),
                total_amount=_to_int(cells[10]),
                national_pension=_to_int(cells[11]),
                health_insurance=_to_int(cells[12]),
                employment_insurance=_to_int(cells[13]),
                longterm_care=_to_int(cells[14]),
                income_tax=_to_int(cells[15]),
                local_tax=_to_int(cells[16]),
                student_loan=_to_int(cells[17]),
                settlement_insurance=_to_int(cells[18]),
                rent_support=_to_int(cells[19]),
                deduction_total=_to_int(cells[20]),
                net_payment=_to_int(cells[21]),
            ))
        if parsed:
            logger.info("[wehago-parser] parsed %d rows deterministically", len(parsed))
            return parsed
    return None


__all__ = ["WehagoPayrollRow", "parse_wehago_workbook"]
