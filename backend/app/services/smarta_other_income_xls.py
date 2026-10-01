"""SmartA 기타소득자료입력 일괄등록 엑셀(.xlsx) 생성.

더존 SmartA 10 급여관리 > 기타소득자료입력 메뉴가 Excel DATA를 자동 변환할 때 읽는
서식. 원본 서식(gita_excel_2022.xls, 2026-09-30 사용자가 실제로 받아 컬럼 분석)을
그대로 재현한다.

**2026-10-01 — xlwt(.xls/BIFF) 포기, openpyxl(.xlsx/OOXML)로 전환.** 2026-09-30엔
"엑셀 업로드가 더존 파서 버그로 막힘"으로 결론 내고 그리드 직접입력으로 전환했었다
(plan/16 §13-3-3). 그런데 사용자가 Excel 웹(OneDrive)에서 원본 서식에 **직접 타이핑**해
업로드에 성공한 파일을 분석해보니, 진짜 원인은 "파서 결함"이 아니라 **xlwt가 쓰는
BIFF(.xls) 포맷 자체가 숫자를 항상 배정밀도 실수로 저장**해 주민등록번호 정수를
`7707281323914.0`처럼 깨뜨리는 것이었다(2026-09-30 실기로 겪은 "13자리가 아닙니다"
오류의 진짜 원인). 성공한 파일은 주민등록번호 셀이 **정수값 + 커스텀 숫자서식
("000000\\-0000000")** — 이 조합을 BIFF+xlwt로 시도했을 때("엑셀불러오기 중 문제가
발생하였습니다")는 실패했지만, OOXML(.xlsx)+openpyxl로는 성공했다. 성공한 파일은 또한
거주(E)/내·외국인(F) 컬럼이 "0.거주"/"0.내국인"이 아니라 접두사 없는 "거주"/"내국인"
이었다 — 소득구분(J)에만 "코드.라벨" 접두사 규칙이 적용된다.
"""

from __future__ import annotations

import logging
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.worksheet import Worksheet

from app.models.payroll import PayrollEntry
from app.services.smarta_business_xls import _rrn_digits

logger = logging.getLogger(__name__)


TITLE_ROW = (
    " 기타소득자료입력     ※ 본 서식은 SmartA 10 급여관리 기타소득자료입력 메뉴로 "
    "Excel DATA를 자동 변환하기 위한 서식입니다."
)

GUIDE_ROW = "\n".join([
    "◈ 기타소득자료입력으로 엑셀 데이터를 반영하기 전 기타소득자입력의 사원정보가 있는 경우는 해당소득자로 반영되며, 소득자가 존재하지 않는 경우 자동으로 추가됩니다.",
    "",
    "◈ 귀속월, 지급월일, 소득자명, 거주, 내/외국인, 주민등록번호, 소득구분은 반드시 기재하셔야 합니다. 변환시 기타소득자입력의 주민(외국인)등록번호로 인식하여 변환이 됩니다.",
    "◈ 주민등록번호를 입력할 때에는 ' : / ? -\"등 특수문자를 제외하고 숫자 13자리를 입력하여 주십시오. ",
    "◈ 소득금액은 지급총액에서 필요경비를 차감하여 자동 산정되며, 업로드시에는 지급총액과 필요경비가 업로드 됩니다.",
    "◈ 소득세와 지방소득세는 지급총액에서 필요경비를 차감한 소득금액을 세율에 따라 자동 계산되어 반영되며, 직접 수정 가능합니다.",
    "◈ 세액계는 소득세와 지방소득세 금액을 합산하여 자동반영되며, 직접 수정 가능합니다.",
    "◈ 차인지급액은 소득금액에서 세액계 금액을 차감한 금액으로 자동반영되며, 직접 수정 가능합니다.",
    "",
    "▣ 본서식을 변형하여 사용시 기타소득자료입력 메뉴로 변환되지 않을 수 있습니다.",
])

COLUMNS = [
    "귀속년월", "지급년월일", "소득자명", "거주", "내/외국인", "기본주소", "상세주소",
    "주민등록번호", "소득구분", "영수일자", "지급총액", "필요경비", "소득금액",
    "세율(%)", "소득세", "지방소득세",
]

# 거주자용 소득구분 코드 → 라벨 (원본 서식 col20, 2026-09-30 실측). 비거주자용(col21,
# 40번대)은 아직 안 다룬다 — 필요해지면 여기에 별도 매핑을 추가한다.
RESIDENT_INCOME_CODE_LABELS: dict[str, str] = {
    "60": "필요경비 없는 기타소득",
    "61": "주식매수선택권 행사이익",
    "62": "그 밖에 필요경비 있는 기타소득(64/68/69/71~77/79제외)",
    "63": "소기업소상공인공제부금해지 소득",
    "64": "서화.골동품 양도소득(필요경비80(90)%)",
    "65": "직무발명보상금(비과세한도 초과분)",
    "68": "비과세기타소득",
    "69": "분리과세기타소득",
    "71": "상금 및 부상(필요경비 80%)",
    "72": "광업권 등(필요경비 60%)",
    "73": "지역권 등(필요경비 60%)",
    "74": "주택입주지체상금(필요경비 80%)",
    "75": "원고료 등(필요경비 60%)",
    "76": "강연료 등 (필요경비 60%)",
    "77": "종교인 소득",
    "78": "사례금",
    "79": "자문료 등 (필요경비 60%)",
    "80": "통신판매 대여 소득(필요경비 60%%)",
}

_RESIDENT_LABEL = "거주"
_NATIONAL_LABEL = "내국인"

# 2026-10-01 실측(gita_excel_2022 (2).xlsx, 업로드 성공본) — 주민등록번호 컬럼의
# 실제 숫자서식. 하이픈은 리터럴이라 이스케이프(`\`)가 필요하다.
_RRN_NUMBER_FORMAT = r"000000\-0000000"

_HEADER_ROW = 3  # 1-indexed(openpyxl). 1=제목, 2=안내문, 3=헤더, 4~=데이터


def _yyyymm(d: date) -> str:
    return f"{d.year}.{d.month:02d}"


def _yyyymmdd(d: date) -> str:
    return f"{d.year}.{d.month:02d}.{d.day:02d}"


def _write_header(ws: Worksheet) -> None:
    header_font = Font(bold=True)
    header_align = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin")
    header_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.cell(row=1, column=1, value=TITLE_ROW)
    ws.cell(row=2, column=1, value=GUIDE_ROW)
    for col in range(1, len(COLUMNS) + 2):
        cell = ws.cell(row=_HEADER_ROW, column=col, value="" if col == 1 else COLUMNS[col - 2])
        cell.font = header_font
        cell.alignment = header_align
        cell.border = header_border


def generate_smarta_other_income_xls(entries: list[PayrollEntry], period: str) -> bytes:
    """기타소득 항목을 SmartA 기타소득자료입력 서식(.xlsx)으로 변환.

    Args:
        entries: income_type=OTHER인 PayrollEntry 리스트. employee가 eager-load되어야 함.
        period: "YYYY-MM" — 문서화 목적. 각 행의 귀속년월·지급년월일은
            `entry.payment_date`(없으면 이 period의 1일)로 채운다.

    Note:
        `other_income_code`가 없거나 `RESIDENT_INCOME_CODE_LABELS`에 없는 코드인 항목은
        위하고가 소득구분을 못 알아봐 업로드에서 제외되므로, 여기서도 건너뛰고 로그만
        남긴다 — 호출자가 "소득구분 미입력" 안내를 사용자에게 보여줄 수 있도록.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    _write_header(ws)

    year, month = period.split("-")
    default_date = date(int(year), int(month), 1)

    row = _HEADER_ROW + 1
    written = 0
    skipped_no_code = 0
    for entry in entries:
        emp = entry.employee
        if not emp:
            continue
        code = entry.other_income_code
        if not code or code not in RESIDENT_INCOME_CODE_LABELS:
            skipped_no_code += 1
            logger.warning(
                "기타소득 소득구분코드 미입력/미지원 — 업로드에서 제외: entry_id=%s code=%s",
                entry.id, code,
            )
            continue

        pay_date = entry.payment_date or default_date

        ws.cell(row=row, column=1, value=written + 1)
        ws.cell(row=row, column=2, value=_yyyymm(pay_date)).number_format = "@"
        ws.cell(row=row, column=3, value=_yyyymmdd(pay_date)).number_format = "@"
        ws.cell(row=row, column=4, value=emp.name)
        ws.cell(row=row, column=5, value=_RESIDENT_LABEL)
        ws.cell(row=row, column=6, value=_NATIONAL_LABEL)
        ws.cell(row=row, column=7, value="")  # 기본주소
        ws.cell(row=row, column=8, value="")  # 상세주소
        rrn = _rrn_digits(emp.rrn_encrypted)
        rrn_cell = ws.cell(row=row, column=9, value=int(rrn) if rrn else "")
        rrn_cell.number_format = _RRN_NUMBER_FORMAT
        ws.cell(row=row, column=10, value=f"{code}.{RESIDENT_INCOME_CODE_LABELS[code]}")  # 소득구분
        ws.cell(row=row, column=11, value=_yyyymmdd(pay_date)).number_format = "@"  # 영수일자
        ws.cell(row=row, column=12, value=entry.total_amount).number_format = "#,##0"  # 지급총액
        ws.cell(row=row, column=13, value=entry.necessary_expense or 0).number_format = "#,##0"  # 필요경비
        ws.cell(row=row, column=14, value="")  # 소득금액 — 자동계산
        ws.cell(row=row, column=15, value="")  # 세율(%) — 자동계산
        ws.cell(row=row, column=16, value=entry.income_tax if entry.income_tax is not None else "")
        ws.cell(row=row, column=17, value=entry.local_tax if entry.local_tax is not None else "")
        row += 1
        written += 1

    if skipped_no_code:
        logger.warning(
            "기타소득 엑셀 생성 — 소득구분코드 미입력/미지원으로 %d명 제외 (총 %d명 중)",
            skipped_no_code, len(entries),
        )

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


__all__ = ["generate_smarta_other_income_xls", "RESIDENT_INCOME_CODE_LABELS"]
