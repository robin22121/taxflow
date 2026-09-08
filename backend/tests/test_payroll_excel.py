"""급여대장 엑셀(위하고T 업로드 양식) 생성기 테스트 (in-memory, no DB).

기준 양식: (임시)주식회사 동문-202511.xlsx — 22컬럼, 2행 병합 헤더, 하단 합계행.
"""

from datetime import date
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.services.payroll_excel import PayrollExcelError, generate_payroll_excel


def _entry(name, *, code="1", bonus=0, meal=200_000, car=200_000, childcare=0,
           income_tax=22_740, local_tax=2_270, nps=91_440, hi=74_440,
           ltc=9_640, ei=0, total=2_500_000):
    return SimpleNamespace(
        raw_name=name,
        employee=SimpleNamespace(name=name, employee_code=code),
        payment_date=date(2025, 11, 25),
        bonus_amount=bonus,
        meal_amount=meal,
        car_amount=car,
        childcare_amount=childcare,
        income_tax=income_tax,
        local_tax=local_tax,
        national_pension=nps,
        health_insurance=hi,
        longterm_care=ltc,
        employment_insurance=ei,
        total_amount=total,
    )


def _sheet(entries, period="2025-11", client_name="주식회사 동문"):
    return load_workbook(BytesIO(
        generate_payroll_excel(entries, period, client_name)
    )).active


def test_header_matches_wehago_template():
    ws = _sheet([_entry("백재봉")])

    assert ws.max_column == 22
    assert [ws.cell(1, c).value for c in (1, 2, 3, 4, 5, 6, 12, 22)] == [
        "사원코드", "사원명", "부서", "직급", "직종", "수당", "공제", "차인지급액",
    ]
    assert [ws.cell(2, c).value for c in range(6, 22)] == [
        "기본급", "상여", "식대", "자가운전", "육아", "지급액계",
        "국민연금", "건강보험", "고용보험", "장기요양보험료",
        "소득세", "지방소득세", "학자금상환액", "정산보험료", "월세지원금", "공제액계",
    ]

    merged = {str(r) for r in ws.merged_cells.ranges}
    assert {"A1:A2", "B1:B2", "C1:C2", "D1:D2", "E1:E2",
            "F1:K1", "L1:U1", "V1:V2"} <= merged


def test_data_row_amounts():
    ws = _sheet([_entry("백재봉")])

    row = [ws.cell(3, c).value for c in range(1, 23)]
    # 부서·직급·직종은 Employee에 대응 필드가 없어 빈 칸
    assert row[:5] == ["1", "백재봉", None, None, None]
    # 기본급 = 2,500,000 − 상여 0 − 식대 200,000 − 자가운전 200,000
    assert row[5:11] == [2_100_000, 0, 200_000, 200_000, 0, 2_500_000]
    # 공제: 4대보험 + 세금 + 학자금/정산보험료/월세지원금(0) + 공제액계
    assert row[11:21] == [91_440, 74_440, 0, 9_640, 22_740, 2_270, 0, 0, 0, 200_530]
    assert row[21] == 2_500_000 - 200_530


def test_negative_income_tax_is_preserved():
    """중도정산 환급으로 소득세가 음수인 경우도 그대로 반영된다."""
    ws = _sheet([_entry("서희석", total=416_660, meal=33_330, car=33_330,
                        income_tax=-76_820, local_tax=-7_640, nps=91_440,
                        hi=12_400, ltc=1_600, ei=3_150)])

    row = [ws.cell(3, c).value for c in range(1, 23)]
    assert row[15] == -76_820 and row[16] == -7_640          # P: 소득세, Q: 지방소득세
    assert row[20] == 91_440 + 12_400 + 3_150 + 1_600 - 76_820 - 7_640  # U: 공제액계
    assert row[21] == 416_660 - row[20]                      # V: 차인지급액


def test_summary_row_is_computed_values():
    ws = _sheet([_entry("백재봉"), _entry("김영희", code="3", ei=18_900)])

    sum_row = 5  # 헤더 2행 + 데이터 2행
    assert ws.cell(sum_row, 1).value == "합계"
    assert f"A{sum_row}:E{sum_row}" in {str(r) for r in ws.merged_cells.ranges}
    # 수식이 아닌 계산된 정수 (위하고T 파서가 수식 캐시를 읽지 못하는 경우 대비)
    assert ws.cell(sum_row, 6).value == 2_100_000 * 2
    assert ws.cell(sum_row, 11).value == 5_000_000
    assert ws.cell(sum_row, 14).value == 18_900


def test_empty_and_bad_period_raise():
    with pytest.raises(PayrollExcelError):
        generate_payroll_excel([], "2025-11")
    with pytest.raises(PayrollExcelError):
        generate_payroll_excel([_entry("백재봉")], "2025/11")
