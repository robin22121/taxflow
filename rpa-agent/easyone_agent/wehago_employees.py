"""위하고 사원등록 [사원자료 엑셀변환] 파일 → 서버로 보낼 사원 기본사항.

파일은 49열(사원코드·사원명·주민번호·입퇴사일·주소·보험료·계좌·연락처 …)이다.
여기서 **필요한 8개 항목만** 뽑는다 — 계좌·연락처·보험료·주소는 노트북 밖으로 보내지 않는다
(plan/16 §12, 10-privacy-security.md). 헤더 이름으로 열을 찾는다.
"""

from __future__ import annotations

import re
import warnings
from datetime import date
from pathlib import Path
from typing import Any

# 위하고 헤더 → 서버 필드
_COLUMNS = {
    "사원코드": "employee_code",
    "사원명": "name",
    "주민(외국인)등록번호": "rrn",
    "입사일자": "hired_at",
    "퇴사일자": "resigned_at",
    "부서": "department",
    "직급": "position",
    "직종": "job_type",
}
_REQUIRED = ("사원코드", "사원명")
# 위하고가 빈 값 대신 넣는 표시
_PLACEHOLDERS = {"직급없음", "--", "-"}


class EmployeeExportError(ValueError):
    """사원자료 엑셀 형식이 예상과 다르다 — 위하고 화면·양식이 바뀌었을 수 있다."""


def _text(value: Any) -> str | None:
    if value is None:
        return None
    s = str(value).strip()
    return None if not s or s in _PLACEHOLDERS else s


def _date(value: Any) -> str | None:
    """YYYYMMDD / YYYY-MM-DD / date → ISO 문자열."""
    if isinstance(value, date):
        return value.isoformat()
    digits = re.sub(r"\D", "", str(value or ""))
    if len(digits) != 8:
        return None
    try:
        return date(int(digits[:4]), int(digits[4:6]), int(digits[6:])).isoformat()
    except ValueError:
        return None


def parse_employee_export(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    with warnings.catch_warnings():  # 위하고 파일에 기본 스타일이 없어 경고가 난다
        warnings.simplefilter("ignore")
        wb = load_workbook(path, read_only=True, data_only=True)
    try:
        # read_only 는 파일을 열어 둔다 — 닫지 않으면 Windows 에서 호출자가 파일을 지우지 못한다
        rows = list(wb.worksheets[0].iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        raise EmployeeExportError("사원자료 엑셀이 비어 있습니다")
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    missing = [h for h in _REQUIRED if h not in header]
    if missing:
        raise EmployeeExportError(f"사원자료 엑셀에 {', '.join(missing)} 열이 없습니다")
    index = {field: header.index(col) for col, field in _COLUMNS.items() if col in header}

    employees = []
    for row in rows[1:]:
        values = {field: row[i] if i < len(row) else None for field, i in index.items()}
        code, name = _text(values.get("employee_code")), _text(values.get("name"))
        if not code or not name:
            continue  # 빈 행
        employees.append({
            "employee_code": code,
            "name": name,
            "rrn": _text(values.get("rrn")),
            "hired_at": _date(values.get("hired_at")),
            "resigned_at": _date(values.get("resigned_at")),
            "department": _text(values.get("department")),
            "position": _text(values.get("position")),
            "job_type": _text(values.get("job_type")),
        })
    return employees
