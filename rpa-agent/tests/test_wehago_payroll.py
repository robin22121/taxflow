from __future__ import annotations

from easyone_agent.wehago import (
    _excel_totals,
    _grid_total,
    _nontaxable_limits,
    _prepare_upload_xlsx,
    _unmapped_amount_columns,
)


def _col(excel: str, wehago: str, has_amount: bool) -> dict:
    return {"col": "?", "excel": excel, "wehago": wehago, "has_amount": has_amount}


def test_unmapped_column_with_amount_is_reported():
    # 2026-09-25 실측: 서도(더미)에 식대 수당이 없어 ③ 매핑이 빈칸 → 식대 200,000원이 빠질 뻔함
    mapping = [
        _col("사원코드", "사원번호(Code)", True),
        _col("기본급", "기본급", True),
        _col("식대", "", True),
        _col("월세지원금", "", False),  # 금액 없음 — 연결 안 돼도 손실 없음
        _col("부서", "", False),  # 문자 열
    ]
    assert _unmapped_amount_columns(mapping) == ["식대"]


def test_employee_code_and_name_must_be_mapped():
    mapping = [_col("사원코드", "", True), _col("사원명", "", False), _col("부서", "", False)]
    assert _unmapped_amount_columns(mapping) == ["사원코드", "사원명"]


def test_total_columns_need_no_mapping():
    mapping = [_col(name, "", True) for name in ("지급액계", "공제액계", "차인지급액")]
    assert _unmapped_amount_columns(mapping) == []


def _payroll_xlsx(path, childcare: int) -> None:
    """이지원천 업로드 양식 (2단 헤더, F~K 수당) 최소 재현 — 사원 1명."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["사원코드", "사원명", "부서", "직급", "직종", "수당"])
    for col in "ABCDE":
        ws.merge_cells(f"{col}1:{col}2")
    ws.merge_cells("F1:K1")
    ws["F2"], ws["G2"], ws["H2"], ws["I2"], ws["J2"], ws["K2"] = "기본급", "상여", "식대", "자가운전", "보육수당", "지급액계"
    ws.append(["1", "한지민", "", "", "", 3_000_000, 0, 200_000, 0, childcare, 3_200_000 + childcare])
    wb.save(path)


def _headers(path) -> list:
    from openpyxl import load_workbook

    return [c.value for c in load_workbook(path).active[1][7:10]]  # 제목 한 줄 양식의 H~J


# 서도(더미) 위하고 실측 수당 목록 — 보육수당이 "육아수당"으로 등록돼 있음
SEODO_ALLOWANCES = [
    {"nm_allow": "기본급", "cd_freeref": None},
    {"nm_allow": "식대", "cd_freeref": "P01"},
    {"nm_allow": "자가운전", "cd_freeref": "H03"},
    {"nm_allow": "육아수당", "cd_freeref": "Q02"},
]


def test_nontaxable_headers_renamed_to_wehago_names(tmp_path):
    src, dst = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
    _payroll_xlsx(src, childcare=200_000)
    assert _prepare_upload_xlsx(src, dst, SEODO_ALLOWANCES) == []
    assert _headers(dst) == ["식대", "자가운전", "육아수당"]


def test_missing_nontaxable_allowance_without_amount_is_ok(tmp_path):
    # 자가운전·보육수당이 위하고에 없어도 금액이 0이면 올려도 빠지는 돈이 없다
    src, dst = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
    _payroll_xlsx(src, childcare=0)
    only_meal = [{"nm_allow": "식대비", "cd_freeref": "P01"}]
    assert _prepare_upload_xlsx(src, dst, only_meal) == []
    assert _headers(dst) == ["식대비", "자가운전", "보육수당"]


def test_missing_nontaxable_allowance_with_amount_is_reported(tmp_path):
    src, dst = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
    _payroll_xlsx(src, childcare=200_000)
    no_childcare = [a for a in SEODO_ALLOWANCES if a["cd_freeref"] != "Q02"]
    assert _prepare_upload_xlsx(src, dst, no_childcare) == ["보육수당(Q02)"]
    assert not dst.exists()


def test_prepared_upload_has_single_header_row(tmp_path):
    # 2단 헤더 1·2행을 제목으로 고르면 위하고가 사원 두 줄을 한 명으로 읽어 금액이 밀린다 (2026-09-25 사고)
    src, dst = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
    _payroll_xlsx(src, childcare=200_000)
    assert _prepare_upload_xlsx(src, dst, SEODO_ALLOWANCES) == []
    from openpyxl import load_workbook

    ws = load_workbook(dst).active
    assert not ws.merged_cells.ranges
    assert [c.value for c in ws[1][:6]] == ["사원코드", "사원명", "부서", "직급", "직종", "기본급"]
    assert [c.value for c in ws[2][:2]] == ["1", "한지민"]


def test_excel_totals(tmp_path):
    src = tmp_path / "in.xlsx"
    _payroll_xlsx(src, childcare=200_000)
    totals = _excel_totals(src)
    assert (totals.headcount, totals.gross, totals.nontaxable) == (1, 3_400_000, 400_000)


def test_grid_total_splits_nontaxable():
    rows = [
        {"nm_allow": "기본급", "fg_tax": "1", "am_tax": 5_400_000},
        {"nm_allow": "식대", "fg_tax": "2", "am_tax": 400_000},
        {"nm_allow": "자가운전", "fg_tax": "2", "am_tax": 200_000},
    ]
    assert _grid_total(rows) == 6_000_000
    assert _grid_total(rows, nontaxable_only=True) == 600_000


def test_excel_totals_caps_nontaxable_at_wehago_limit(tmp_path):
    # 위하고는 한도(20만) 초과분을 과세로 나눈다 — 보육수당 250,000 → 비과세 200,000
    src = tmp_path / "in.xlsx"
    _payroll_xlsx(src, childcare=250_000)
    limits = _nontaxable_limits([
        {"nm_allow": "식대", "cd_freeref": "P01", "am_tflimit": 200_000},
        {"nm_allow": "육아수당", "cd_freeref": "Q02", "am_tflimit": 200_000},
        {"nm_allow": "기본급", "cd_freeref": None, "am_tflimit": 0},
    ])
    assert limits == {"P01": 200_000, "Q02": 200_000}
    totals = _excel_totals(src, limits)
    assert (totals.gross, totals.nontaxable) == (3_450_000, 400_000)
